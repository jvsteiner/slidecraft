"""
DataTableLayout — data table with optional sidebar and footer.

Content zones:
    title, subtitle, table (list of lists), col_widths (optional),
    sidebar (optional: {title, items}), footer (optional),
    highlight_row (optional int)
"""
from .base import BaseLayout


class DataTableLayout(BaseLayout):
    name = 'data_table'

    def render(self, db, content: dict, slide_num: int):
        slide = db.new_slide()
        p = db.p
        m = db.theme.margin
        cw = self._content_width(db)

        y = p.slide_header(slide, slide_num,
                           content.get('title', ''),
                           content.get('subtitle'))

        table_data = content.get('table', [])
        sidebar = content.get('sidebar')
        footer = content.get('footer')
        highlight_row = content.get('highlight_row')
        col_widths_raw = content.get('col_widths')

        if not table_data:
            return

        # Determine table width based on sidebar presence
        gutter = db.theme.gutter
        if sidebar:
            table_width = cw * 0.6 - gutter / 2
            sidebar_left = m + table_width + gutter
            sidebar_width = cw - table_width - gutter
        else:
            table_width = cw

        # Scale col_widths to fit table_width
        if col_widths_raw:
            total = sum(col_widths_raw)
            col_widths = [w / total * table_width for w in col_widths_raw]
        else:
            col_widths = None

        table_top = y + 0.2
        available_height = 7.5 - table_top - (0.6 if footer else 0.3)
        table_height = min(0.4 * len(table_data), available_height)

        def text_color_fn(row_idx, col_idx, val):
            if row_idx == 0:
                return 'primary'
            if highlight_row is not None and row_idx == highlight_row:
                return 'primary'
            return 'text_secondary'

        def align_fn(row_idx, col_idx):
            if col_idx == 0:
                return 'left'
            return 'center'

        font_size = content.get('font_size', 16)
        tbl_shape = p.table(slide, m, table_top, table_width, table_height,
                            table_data, col_widths=col_widths,
                            header_color='primary',
                            text_color_fn=text_color_fn,
                            align_fn=align_fn,
                            font_size=font_size)

        # Bold + larger font for highlight row
        if highlight_row is not None:
            from pptx.util import Pt
            tbl = tbl_shape.table
            if highlight_row < len(table_data):
                for col_idx in range(len(table_data[0])):
                    cell = tbl.cell(highlight_row, col_idx)
                    for para in cell.text_frame.paragraphs:
                        para.font.bold = True
                        para.font.size = Pt(18)

        # Sidebar
        if sidebar:
            sidebar_top = table_top
            sidebar_height = table_height
            p.rect(slide, sidebar_left, sidebar_top,
                   sidebar_width, sidebar_height,
                   fill='surface', border='border')
            p.bullets(slide, sidebar_left + 0.2, sidebar_top + 0.15,
                      sidebar_width - 0.4, sidebar_height - 0.3,
                      items=sidebar.get('items', []),
                      title=sidebar.get('title'),
                      title_size=18, title_color='primary',
                      bullet_size=font_size, bullet_color='text_secondary')

        # Footer
        if footer:
            footer_y = db.theme.spacing('footer_y') - 0.3
            p.text_box(slide, m, footer_y, cw, 0.4,
                       footer, size=14, color='text_dim')
